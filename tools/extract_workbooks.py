"""Turn the cleaned templates in public/templates/ into src/data/workbooks.json, which
the site uses to draw every sheet, take input and rerun the formulas.

    python tools/extract_workbooks.py
"""
import datetime, json, re
from pathlib import Path
import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

ROOT = Path(__file__).resolve().parent.parent
MARKS = {"4.1": 15, "4.2": 15, "4.3": 10, "4.4": 10}
INPUT_FILL = "FFFBE6"
EPOCH = datetime.datetime(1899, 12, 30)

def serial(d):
    if isinstance(d, datetime.date) and not isinstance(d, datetime.datetime):
        d = datetime.datetime.combine(d, datetime.time())
    return (d - EPOCH).days

def rgb(color):
    if color is None or color.type != "rgb" or not isinstance(color.rgb, str):
        return None
    return color.rgb[-6:]

def validations(ws):
    """cell -> {t, o} from the sheet's data validations."""
    out = {}
    for dv in ws.data_validations.dataValidation:
        kind = {"decimal": "num", "whole": "int", "list": "list", "date": "date"}.get(dv.type)
        if not kind:
            continue
        spec = {"t": kind}
        if kind == "list":
            spec["o"] = [s.strip() for s in dv.formula1.strip('"').split(",")]
        for rng in str(dv.sqref).split():
            a, _, b = rng.partition(":")
            b = b or a
            c1, r1 = re.match(r"([A-Z]+)(\d+)", a).groups()
            c2, r2 = re.match(r"([A-Z]+)(\d+)", b).groups()
            for r in range(int(r1), int(r2) + 1):
                for c in range(column_index_from_string(c1), column_index_from_string(c2) + 1):
                    out[(r, c)] = spec
    return out

def charts(ws):
    out = []
    for ch in ws._charts:
        series = []
        for s in ch.series:
            name = s.tx.v if s.tx is not None and s.tx.v else (s.tx.strRef.f if s.tx is not None and s.tx.strRef else "")
            series.append({"name": name, "ref": s.val.numRef.f})
        title = None
        if ch.title is not None and ch.title.tx is not None and ch.title.tx.rich is not None:
            title = "".join(r.t for p in ch.title.tx.rich.p for r in (p.r or []))
        out.append({"r1": ch.anchor._from.row + 1, "r2": ch.anchor.to.row + 1, "title": title, "series": series})
    return out

def extract(path: Path):
    wb = openpyxl.load_workbook(path)
    styles, fmts = [], []
    def idx(lst, item):
        if item not in lst:
            lst.append(item)
        return lst.index(item)
    sheets = []
    for ws in wb.worksheets:
        dvs = validations(ws)
        merged = {}
        merges = []
        for m in ws.merged_cells.ranges:
            merges.append([m.min_row, m.min_col, m.max_row, m.max_col])
            for r in range(m.min_row, m.max_row + 1):
                for c in range(m.min_col, m.max_col + 1):
                    if (r, c) != (m.min_row, m.min_col):
                        merged[(r, c)] = True
        is_cover = ws.title.startswith("1.")
        cells = []
        for row in ws.iter_rows():
            for c in row:
                if (c.row, c.column) in merged:
                    continue
                fill = rgb(c.fill.fgColor) if c.fill is not None and c.fill.fill_type == "solid" else None
                v = c.value
                if v is None and fill is None:
                    continue
                f = None
                if isinstance(v, str) and v.startswith("="):
                    f, v = v[1:], None
                elif isinstance(v, (datetime.date, datetime.datetime)):
                    v = serial(v)
                st = {"bg": fill, "fg": rgb(c.font.color) if c.font and c.font.color else None,
                      "b": bool(c.font.b), "i": bool(c.font.i), "sz": c.font.sz,
                      "al": c.alignment.horizontal, "wr": bool(c.alignment.wrap_text)}
                cell = [c.row, c.column, v, f, idx(styles, st), idx(fmts, c.number_format)]
                if fill == INPUT_FILL and not is_cover and f is None:
                    spec = dvs.get((c.row, c.column))
                    if spec is None:
                        fmt = c.number_format
                        spec = {"t": "date" if "yy" in fmt else "num" if re.search(r"0\.0|#,##0", fmt) else "text"}
                    cell.append(spec)
                cells.append(cell)
        widths = {k: round(d.width, 1) for k, d in ws.column_dimensions.items() if d.width}
        sheets.append({"name": ws.title, "widths": widths, "merges": merges, "charts": charts(ws),
                       "maxCol": ws.max_column, "maxRow": ws.max_row, "cells": cells})
    cover = wb.worksheets[0]
    mid = re.match(r"NAAC_Metric_(\d\.\d)", path.name).group(1)
    title = re.sub(r"\s*\(Weightage.*", "", cover["B6"].value).split("–", 1)[-1].strip()
    return {"id": mid, "title": title, "marks": MARKS[mid], "file": "templates/" + path.name,
            "styles": styles, "fmts": fmts, "sheets": sheets}

if __name__ == "__main__":
    books = [extract(p) for p in sorted((ROOT / "public" / "templates").glob("NAAC_Metric_*.xlsx"))]
    js = json.dumps(books, ensure_ascii=False, separators=(",", ":"))
    out = ROOT / "src" / "data" / "workbooks.json"
    out.write_text(js, encoding="utf-8")
    for b in books:
        n = sum(1 for s in b["sheets"] for c in s["cells"] if len(c) > 6)
        print(b["id"], b["title"], "| sheets", len(b["sheets"]), "| inputs", n, "| charts", [c["title"] for s in b["sheets"] for c in s["charts"]])
    print(out, out.stat().st_size, "bytes")
