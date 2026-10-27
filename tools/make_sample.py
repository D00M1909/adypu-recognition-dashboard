"""Build src/data/sample.json: illustrative data for every input sheet of the
four NAAC Attribute 4 workbooks, so the site can show what a filled
workbook looks like. Head totals follow each workbook's own worked example
(and the Attribute 4 booklet), spread across the sub-heads.

    python tools/make_sample.py
"""
import datetime, json
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
FY = [4, 5, 6, 7]                      # columns D-G: FY 2025-26 back to FY 2022-23
EPOCH = datetime.date(1899, 12, 30)
serial = lambda y, m, d: (datetime.date(y, m, d) - EPOCH).days

class Book:
    def __init__(self, name):
        self.path = next((ROOT / "public" / "templates").glob(name))
        self.wb = openpyxl.load_workbook(self.path)
        self.names = [ws.title for ws in self.wb.worksheets]
        self.values = {}
    def ws(self, sheet):
        return self.wb[sheet], self.names.index(sheet)
    def row(self, sheet, code, first=False):
        """Row whose column-B code is `code`; `first` takes the top one when a later table reuses codes."""
        ws, _ = self.ws(sheet)
        hits = [r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 2).value or "").strip() == code]
        assert hits and (first or len(hits) == 1), (sheet, code, hits)
        return hits[0]
    def put(self, sheet, r, c, v):
        ws, si = self.ws(sheet)
        cell = ws.cell(r, c)
        assert cell.fill.fill_type == "solid" and cell.fill.fgColor.rgb == "FFFFFBE6", (sheet, cell.coordinate, "not an input")
        self.values[f"{si}!{r},{c}"] = v
    def years(self, sheet, code, vals, col0=4, first=False):
        r = self.row(sheet, code, first)
        for i, v in enumerate(vals):
            if v is not None:
                self.put(sheet, r, col0 + i, v)
        return r

def split(totals, weights):
    """Spread each FY total over sub-heads by weight; the first sub-head takes the rounding."""
    out = [[round(t * w / sum(weights), 2) for t in totals] for w in weights]
    for i, t in enumerate(totals):
        out[0][i] = round(t - sum(o[i] for o in out[1:]), 2)
    return out

def heads(b, sheet, plan, ref):
    for n, (codes, totals, weights) in enumerate(plan):
        for code, vals in zip(codes, split(totals, weights)):
            r = b.years(sheet, code, [v or None for v in vals])
            if any(vals):
                b.put(sheet, r, 10, f"Sch. {ref + n}, p. {12 + 2 * n}")

def checklist(b, sheet, rows, na=()):
    ws, _ = b.ws(sheet)
    for k in range(1, rows + 1):
        r = b.row(sheet, str(k))
        for i, c in enumerate(FY):
            status = "Not Applicable" if k in na else "Pending" if (i == 0 and k % 4 == 0) else "Available"
            b.put(sheet, r, c + 1, status)          # status columns are E-H

# ---------- 4.1 Capital income vs expenditure ----------
def m41():
    b = Book("NAAC_Metric_4.1_*.xlsx")
    heads(b, "3. Capital Income", [
        (["A1", "A2", "A3", "A4"], [150, 120, 80, 60], [4, 2, 2.5, 1.5]),
        (["B1", "B2"], [900, 800, 700, 600], [8, 2]),
        (["C1", "C2"], [500, 450, 300, 250], [7, 3]),
        (["D1", "D2", "D3"], [60, 40, 35, 30], [3, 3, 4]),
        (["E1"], [0, 10, 5, 0], [1]),
    ], 3)
    heads(b, "4. Capital Expenditure", [
        (["1.1", "1.2"], [600, 700, 380, 300], [7, 3]),
        (["2.1", "2.2", "2.3", "2.4"], [140, 120, 100, 80], [3, 3, 2.5, 1.5]),
        (["3.1", "3.2"], [300, 250, 200, 150], [8, 2]),
        (["4.1", "4.2", "4.3", "4.4"], [350, 280, 220, 190], [4, 3, 1, 2]),
        (["5.1"], [0, 40, 20, 0], [1]),
        (["6.1", "6.2", "6.3", "6.4"], [120, 100, 90, 70], [3, 3, 2.5, 1.5]),
        (["7.1"], [45, 40, 35, 30], [1]),
        (["8.1"], [5, 0, 0, 10], [1]),
    ], 8)
    checklist(b, "7. Document Checklist", 12, na={10})
    return b

# ---------- 4.2 Revenue income vs expenditure ----------
INC42 = [
    (["A1", "A2", "A3", "A4"], [9500, 8800, 7600, 6800], [70, 10, 6, 14]),
    (["B1", "B2"], [250, 200, 150, 120], [6, 4]),
    (["D1"], [20, 15, 10, 10], [1]),
    (["E1", "E2"], [80, 60, 50, 40], [6, 4]),
    (["F1"], [30, 25, 20, 20], [1]),
    (["G1", "G2"], [600, 550, 500, 450], [8, 2]),
    (["H1", "H2", "H3"], [220, 180, 150, 120], [6, 2, 2]),
]
EXP42 = [
    (["1.1", "1.2"], [5200, 4800, 4400, 4000], [7, 3]),
    (["2.1", "2.2", "2.3"], [700, 650, 600, 550], [4, 3, 3]),
    (["3.1", "3.2", "3.3"], [650, 600, 560, 520], [4, 3, 3]),
    (["4.1", "4.2"], [500, 450, 400, 380], [6, 4]),
    (["5.1"], [300, 280, 260, 240], [1]),
    (["6.1"], [250, 230, 210, 200], [1]),
    (["7.1", "7.2"], [120, 100, 90, 80], [6, 4]),
    (["8.1", "8.2"], [600, 550, 500, 450], [7, 3]),
    (["9.1", "9.2"], [180, 160, 140, 120], [7, 3]),
    (["10.1", "10.2", "10.3", "10.4", "10.5"], [200, 180, 170, 150], [3, 1, 3, 2, 1]),
]
def m42():
    b = Book("NAAC_Metric_4.2_*.xlsx")
    heads(b, "3. Revenue Income", INC42, 11)
    heads(b, "4. Revenue Expenditure", EXP42, 18)
    inc = [sum(p[1][i] for p in INC42) for i in range(4)]
    exp = [sum(p[1][i] for p in EXP42) for i in range(4)]
    dep, prov, excl = [820, 760, 690, 640], [40, 35, 30, 25], [35, 30, 25, 20]
    s = "5. Summary & Ratio"
    b.years(s, "E1", [i + x for i, x in zip(inc, excl)])
    b.years(s, "E2", excl)
    b.years(s, "E4", [e + d + p for e, d, p in zip(exp, dep, prov)])
    b.years(s, "E5", dep)
    b.years(s, "E6", prov)
    checklist(b, "7. Document Checklist", 11)
    return b

# ---------- 4.3 Financial sustainability ----------
def m43():
    b = Book("NAAC_Metric_4.3_*.xlsx")
    s = "3. Corpus Fund"
    b.put(s, b.row(s, "A"), 7, 2000)                               # opening corpus, FY 2022-23 only
    adds = {"B1": [1200, 1000, 700, 500], "B2": [300, 250, 200, 200], "B3": [100, 80, 60, 50], "B4": [380, 280, 200, 140]}
    for code, vals in adds.items():
        b.years(s, code, vals)
    b.years(s, "C1", [120, 90, 60, None])
    b.years(s, "C2", [80, 60, 40, None])
    closing, c = [], 2000
    for i in (3, 2, 1, 0):
        c += sum(v[i] for v in adds.values()) - [200, 150, 100, 0][i]
        closing.insert(0, c)
    b.years(s, "E1", closing)                                      # audited figure matches: 7,190 at 31 Mar 2026
    b.years(s, "E4", [8700, 8000, 7330, 6690])
    s = "4. Investment Portfolio"
    for code, vals in {"1": [4500, 4000, 3500, 3200], "2": [1500, 1100, 600, 200], "3": [800, 600, 300, None],
                       "4": [900, 500, 200, None], "5": [300, 200, 100, None]}.items():
        b.years(s, code, vals, first=True)
    b.years(s, "S1", [6, 5, 4, 3])
    b.years(s, "S2", [560, 450, 340, 250])
    b.years(s, "S3", [8000, 6400, 4700, 3400])
    s = "5. Strategies & Policies"
    rows = [
        ("Corpus Fund Rules adopted; principal preserved, only income used for scholarships and capital works.", "Board of Management", "14-Jun-2022", "Implemented"),
        ("Not less than 10% of the annual operating surplus moved to the corpus by resolution.", "Finance Committee", "22-Jul-2022", "Implemented"),
        ("Sponsoring trust commits a yearly contribution to the corpus.", "Governing Body", "05-Apr-2022", "Implemented"),
        ("Alumni and CSR endowments for chairs and merit scholarships.", "Board of Management", "18-Nov-2022", "Implemented"),
        ("Endowment fund kept as required by the State Private Universities Act.", "Governing Body", "05-Apr-2022", "Implemented"),
        ("Investment Policy: permitted instruments, 40% cap per bank, AA+ minimum rating.", "Finance Committee", "09-Sep-2023", "Implemented"),
        ("Funds spread over six banks and five instrument types.", "Finance Committee", "09-Sep-2023", "Implemented"),
        ("Investment Committee reviews the portfolio every quarter.", "Finance Committee", "09-Sep-2023", "Implemented"),
        ("Consultancy, executive education and research overheads grown to cut fee dependency.", "Board of Management", "11-Jan-2024", "In progress"),
        ("Central procurement and rooftop solar to cut running costs.", "Finance Committee", "11-Jan-2024", "Implemented"),
        ("Contingency fund for loan repayment and major repairs.", "", "", "Planned"),
    ]
    for k, (desc, body, date, status) in enumerate(rows, 1):
        r = b.row(s, str(k))
        b.put(s, r, 4, desc)
        if body: b.put(s, r, 5, body); b.put(s, r, 6, date); b.put(s, r, 8, f"Resolution file FIN/{2022 + k % 3}/{k:02d}")
        b.put(s, r, 7, status)
    checklist(b, "8. Document Checklist", 11)
    return b

# ---------- 4.4 Financial controls & risk management ----------
def m44():
    b = Book("NAAC_Metric_4.4_*.xlsx")
    s = "3. Internal Audit"
    b.years(s, "IA1", ["Yes"] * 4)
    b.years(s, "IA2", ["Example & Co., Chartered Accountants (external CA firm)"] * 4)
    b.years(s, "IA3", [serial(y, 4, 10) for y in (2025, 2024, 2023, 2022)])
    b.years(s, "IA4", ["Quarterly", "Quarterly", "Quarterly", "Half-yearly"])
    b.years(s, "IA5", [4, 4, 4, 2])
    b.years(s, "IA6", [38, 45, 52, 40])
    b.years(s, "IA7", [35, 40, 44, 30])
    b.years(s, "IA10", ["Yes", "Yes", "Yes", "No"])
    b.years(s, "IA11", [serial(y, 5, 20) for y in (2026, 2025, 2024, 2023)])
    covered = [12, 11, 10, 8]
    for k in range(1, 13):
        b.years(s, f"C{k}", ["Yes" if k <= n else "No" for n in covered])
    s = "4. Statutory & External Audit"
    b.years(s, "SA1", ["Example Associates, Chartered Accountants"] * 4)
    b.years(s, "SA2", ["123456W"] * 4)
    b.years(s, "SA3", ["Board of Management"] * 4)
    b.years(s, "SA4", [serial(2026, 8, 28), serial(2025, 9, 10), serial(2024, 9, 25), serial(2023, 10, 15)])
    b.years(s, "SA6", ["Unqualified"] * 4)
    b.years(s, "SA7", [6, 8, 10, 12])
    b.years(s, "SA8", [6, 7, 8, 8])
    b.years(s, "SA11", [serial(2026, 9, 12), serial(2025, 9, 24), serial(2024, 10, 8), serial(2023, 10, 30)])
    b.years(s, "SA12", [serial(2026, 9, 26), serial(2025, 10, 6), serial(2024, 10, 22), serial(2023, 11, 14)])
    b.years(s, "SA13", ["Yes"] * 4)
    for code, vals in {"B1": ["Yes"] * 4, "B2": ["Yes"] * 4, "B3": ["Yes", "Yes", "Yes", "No"],
                       "B4": ["Not Applicable"] * 4, "B5": ["Not Applicable"] * 4}.items():
        b.years(s, code, vals)
    s = "5. Risk Mgmt Approaches"
    done, partial = [11, 9, 7, 4], [1, 2, 2, 3]
    for k in range(1, 13):
        b.years(s, str(k), ["Implemented" if k <= d else "Partially implemented" if k <= d + p else "Not implemented"
                            for d, p in zip(done, partial)])
    s = "6. Risk Register"
    scores = [(3, 4, "Mitigating"), (3, 3, "Mitigating"), (4, 3, "Open"), (1, 4, "Closed"), (2, 5, "Mitigating"),
              (2, 3, "Closed"), (2, 4, "Mitigating"), (2, 3, "Open"), (2, 2, "Closed"), (1, 4, "Closed")]
    owners = ["Registrar", "Finance Officer", "Finance Officer", "Finance Committee", "Internal Auditor",
              "Accounts Head", "IT Head", "Finance Officer", "Research Office", "Estate Office"]
    for k, ((l, i, st), owner) in enumerate(zip(scores, owners), 1):
        r = b.row(s, str(k))
        b.put(s, r, 5, l); b.put(s, r, 6, i); b.put(s, r, 10, owner); b.put(s, r, 11, st)
    ws, _ = b.ws(s)
    r = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 3).value or "").startswith("Date of last review"))
    b.put(s, r, 7, serial(2026, 6, 18))
    b.put(s, r, 9, "Reviewed by (Committee): Finance & Audit Committee")
    checklist(b, "9. Document Checklist", 13)
    return b

if __name__ == "__main__":
    books = {"4.1": m41(), "4.2": m42(), "4.3": m43(), "4.4": m44()}
    data = {k: b.values for k, b in books.items()}
    out = ROOT / "src" / "data" / "sample.json"
    out.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    for k, v in data.items():
        print(k, len(v), "cells")
